"""Test workbook update end-to-end, idempotent upsert & validasi."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from tests.conftest import base_cfg, make_event
from src import workbook, validator, downtime as dow, pareto, extractor
from src.issue_mapper import IssueMapper


@pytest.fixture()
def cfg(tmp_path) -> dict:
    cfg = base_cfg()
    cfg["output"] = {
        "template": str(tmp_path / "template" / "Template_Monitoring_Proses2.xlsx"),
        "output_dir": str(tmp_path / "out"),
        "latest_name": "Monitoring_Proses2_LATEST.xlsx",
        "dated_name": "Monitoring_Proses2_{date}.xlsx",
        "backup_dir": str(tmp_path / "out" / "backups"),
        "backup_keep": 10,
        "cache_dir": str(tmp_path / "cache"),
    }
    return cfg


def build_and_validate(cfg, events_for_day, target, prev_events=()):
    merged = extractor.dedup_events(list(prev_events) + list(events_for_day))
    IssueMapper(cfg).apply(merged)
    cur = [e for e in merged if e["date"] == target]
    wb = workbook.load_or_create(cfg)
    workbook.build_sheets(
        wb, cfg, merged=merged, target=target,
        pareto_mesin=pareto.pareto_machine(cur),
        pareto_issue=pareto.pareto_issue(cur),
        trend=dow.trend_rows(merged),
        rekap=dow.rekap_machine_rows(merged, cfg),
        summary=dow.summary(merged, target),
        followup=dow.followup_rows(merged, target, cfg))
    workbook.apply_charts(wb, cfg)
    workbook.finalize(wb, cfg)
    return wb, merged


def test_full_update_26_october(cfg):
    target = dt.date(2026, 10, 26)
    evs = [
        make_event(target, machine="VACUM 7", menit=200, keterangan="alarm produk ganda", kode="BJ-PE00339-0"),
        make_event(target, machine="VACUM 7", menit=20, keterangan="tunggu bahan", kode="BJ-PE00339-0"),
        make_event(target, machine="TFM 1", menit=15, keterangan="produk nempel di plugassit"),
        make_event(target, machine="WM", menit=30, keterangan="bahan lepas dari rantai", kategori="C"),
    ]
    wb, merged = build_and_validate(cfg, evs, target)
    errors = validator.validate_workbook(wb, cfg, target=target)
    assert errors == []

    dated, latest = workbook.save_workbook(wb, cfg, target)
    assert latest.exists() and dated.exists()

    wb2 = openpyxl.load_workbook(latest)
    for sheet in ("RINGKASAN", "DETAIL_DOWNTIME", "REKAP_MESIN", "PARETO_MESIN",
                  "PARETO_ISSUE", "TREND_HARIAN", "PLUG_ASSIST", "TUNGGU_BAHAN", "FOLLOW_UP"):
        assert sheet in wb2.sheetnames
    # pareto data tiambin chart anchored
    assert len(wb2["PARETO_MESIN"]._charts) == 1
    assert len(wb2["TREND_HARIAN"]._charts) == 1


def test_idempotent_same_date_no_duplicates(cfg):
    target = dt.date(2026, 10, 27)
    evs = [
        make_event(target, machine="VACUM 8", menit=120, keterangan="alarm produk ganda"),
        make_event(target, machine="TFM 2", menit=60, keterangan="tunggu bahan"),
    ]
    wb1, _ = build_and_validate(cfg, evs, target)
    workbook.save_workbook(wb1, cfg, target)

    # baca ulang workbook sbg dasar (updapat idempotent)
    prev = workbook.read_events(cfg)
    wb2, merged2 = build_and_validate(cfg, evs, target, prev_events=prev)
    dates27 = [e for e in merged2 if e["date"] == target]
    assert len(dates27) == len(evs)          # tidak berlipat

    workbook.save_workbook(wb2, cfg, target)
    wb3 = openpyxl.load_workbook(Path(cfg["output"]["output_dir"]) / cfg["output"]["latest_name"])
    detail_dates = [r[0] for r in wb3["DETAIL_DOWNTIME"].iter_rows(min_row=3, min_col=1, max_col=1, values_only=True)
                    if r[0] not in (None, "")]
    assert detail_dates.count(dt.datetime(target.year, target.month, target.day)) == len(evs) \
        or detail_dates.count(target) == len(evs)
    errors = validator.validate_workbook(wb3, cfg, target=target)
    assert errors == []


def test_add_second_day_extends_trend_chart(cfg):
    d1, d2 = dt.date(2026, 10, 26), dt.date(2026, 10, 27)
    evs1 = [make_event(d1, machine="TFM 1", menit=50)]
    wb1, _ = build_and_validate(cfg, evs1, d1)
    workbook.save_workbook(wb1, cfg, d1)

    prev = workbook.read_events(cfg)
    evs2 = [make_event(d2, machine="VACUM 9", menit=120), make_event(d2, machine="WM.2", menit=70)]
    wb2, merged = build_and_validate(cfg, evs2, d2, prev_events=prev)
    assert len(merged) == 3
    trend_rows = wb2["TREND_HARIAN"].max_row - 2
    assert trend_rows == 2                    # chart/data range otomatis lonjur
    errors = validator.validate_workbook(wb2, cfg, target=d2)
    assert errors == []


def test_backup_rotate_and_latest(cfg):
    d1, d2 = dt.date(2026, 10, 26), dt.date(2026, 10, 27)
    wb1, _ = build_and_validate(cfg, [make_event(d1, machine="TFM 1", menit=10)], d1)
    workbook.save_workbook(wb1, cfg, d1)
    wb2, _ = build_and_validate(cfg, [make_event(d2, machine="TFM 1", menit=10)], d2)
    workbook.save_workbook(wb2, cfg, d2)
    backups = list(Path(cfg["output"]["backup_dir"]).glob("*.xlsx"))
    assert len(backups) == 1                  # copy LATEST lama sebelum ditimpa
