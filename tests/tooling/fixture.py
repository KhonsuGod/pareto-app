"""Fixture workbook spare sintetis (hanya untuk test, bukan data produksi).

Struktur meniru workbook asli: 3 sheet, header baris 4, formula E/F di
STOCK SPARE, log kosong dengan jangkauan formula $5:$<LOG_LAST>.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl

from src.tooling.repository import LOG_FIRST


def build_spare_workbook(path: Path, routing: list[dict], items: list[dict],
                         log_last: int = 40) -> None:
    wb = openpyxl.Workbook()
    mt = wb.active
    mt.title = "MASTER TOOLING"
    mt["A1"] = "MASTER TOOLING - PROSES 2"
    mt["A3"] = "ROUTING"
    head = ["No.", "Mesin (PPIC)", "Tooling Family", "Kode BJ", "Nama Produk",
            "Cavity PPIC", "Mould", "Insert", "Plug Assist", "Ejector", "Catatan / Exception"]
    for c, h in enumerate(head, start=1):
        mt.cell(4, c).value = h
    for i, r in enumerate(routing, start=5):
        mt.cell(i, 1).value = r.get("no", i - 4)
        mt.cell(i, 2).value = r.get("mesin")
        mt.cell(i, 3).value = r["family"]
        mt.cell(i, 4).value = r.get("kode_bj")
        mt.cell(i, 5).value = r.get("nama_produk")
        mt.cell(i, 6).value = r.get("cavity")
        mt.cell(i, 7).value = r.get("mould", "CEK TS")
        mt.cell(i, 8).value = r.get("insert", "CEK TS")
        mt.cell(i, 9).value = r.get("plug_assist", "CEK TS")
        mt.cell(i, 10).value = r.get("ejector", "CEK TS")
        mt.cell(i, 11).value = r.get("catatan")

    st = wb.create_sheet("STOCK SPARE")
    st["B3"] = "=TODAY()"
    head_st = ["No.", "Tooling Family", "Sparepart / Component", "Stock Awal",
               "Keluar Hari Ini", "Stock Aktual", "Lokasi", "Catatan", "kondisi sparepart"]
    for c, h in enumerate(head_st, start=1):
        st.cell(4, c).value = h
    for i, it in enumerate(items, start=5):
        st.cell(i, 1).value = i - 4
        st.cell(i, 2).value = it["family"]
        st.cell(i, 3).value = it["component"]
        st.cell(i, 4).value = it.get("stock_awal")  # None = belum diinput
        st.cell(i, 5).value = (
            f'=IF(OR(B{i}="",C{i}=""),"",SUMIFS(\'LOG PENGELUARAN\'!$H${LOG_FIRST}:$H${log_last},'
            f"'LOG PENGELUARAN'!$D${LOG_FIRST}:$D${log_last},B{i},"
            f"'LOG PENGELUARAN'!$G${LOG_FIRST}:$G${log_last},C{i},"
            f"'LOG PENGELUARAN'!$A${LOG_FIRST}:$A${log_last},$B$3))")
        st.cell(i, 6).value = (
            f'=IF(D{i}="","",D{i}-SUMIFS(\'LOG PENGELUARAN\'!$H${LOG_FIRST}:$H${log_last},'
            f"'LOG PENGELUARAN'!$D${LOG_FIRST}:$D${log_last},B{i},"
            f"'LOG PENGELUARAN'!$G${LOG_FIRST}:$G${log_last},C{i}))")
        st.cell(i, 8).value = it.get("catatan")

    lg = wb.create_sheet("LOG PENGELUARAN")
    lg["A1"] = "LOG PENGELUARAN SPAREPART | PROSES 2"
    head_log = ["Tanggal", "Jam", "Mesin Tujuan", "Tooling Family", "Kode BJ",
                "Nama Produk", "Sparepart Keluar", "Qty Keluar", "Keterangan / Alasan", "PIC / Checker"]
    for c, h in enumerate(head_log, start=1):
        lg.cell(4, c).value = h
    wb.save(path)


DEFAULT_ROUTING = [
    {"no": 1, "mesin": "TF2", "family": "DATAR-12", "kode_bj": "BJ-PP00526-0",
     "nama_produk": "PLASTIC CUP PP 12 OZ DATAR", "cavity": 12},
    {"no": 2, "mesin": "TF2", "family": "DATAR-12", "kode_bj": "BJ-PP00527-0",
     "nama_produk": "PLASTIC CUP PP 12 OZ DATAR V2", "cavity": 12},
    {"no": 3, "mesin": "WM", "family": "CUP-16", "kode_bj": "BJ-PP00246-0",
     "nama_produk": "CLEAR COLD CUP 16 OZ", "cavity": 15},
    {"no": 4, "mesin": "", "family": "CUP-CORN", "kode_bj": "",
     "nama_produk": "Plastic Cup Corn PP", "cavity": None,
     "catatan": "KODE BJ KOSONG DI PPIC | MESIN KOSONG DI PPIC"},
]

DEFAULT_ITEMS = [
    {"family": "DATAR-12", "component": "MOULD", "stock_awal": None},
    {"family": "DATAR-12", "component": "INSERT", "stock_awal": None},
    {"family": "DATAR-12", "component": "PLUG ASSIST", "stock_awal": None},
    {"family": "DATAR-12", "component": "EJECTOR", "stock_awal": None},
    {"family": "CUP-16", "component": "MOULD", "stock_awal": 5},   # sudah terverifikasi
    {"family": "CUP-16", "component": "INSERT", "stock_awal": 0},  # terverifikasi nol
    {"family": "CUP-CORN", "component": "EJECTOR", "stock_awal": None},
]


def make_store(tmp_path: Path, minimum_stock: dict | None = None):
    from src.tooling.repository import SpareStore
    wb_path = tmp_path / "monitorung_spare_proses2.xlsx"
    build_spare_workbook(wb_path, DEFAULT_ROUTING, DEFAULT_ITEMS)
    cfg = {
        "spare": {
            "workbook": str(wb_path),
            "source_workbook": "",
            "backup_dir": str(tmp_path / "backups"),
            "backup_keep": 3,
            "state_file": str(tmp_path / "state.json"),
            "minimum_stock": minimum_stock or {},
        }
    }
    return SpareStore(cfg), wb_path
